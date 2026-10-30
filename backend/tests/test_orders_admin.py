import io
import json
from pathlib import Path

import pytest
from django.core.management import call_command
from openpyxl import load_workbook

from apps.accounts.roles import Role
from apps.orders.models import Order
from tests.conftest import order_payload, post_order

OWN, A, M, E, N = Role.OWNER, Role.ADMIN, Role.ORDER_MANAGER, Role.CONTENT_EDITOR, Role.ANALYST


@pytest.fixture
def orders(client, sellable):
    sellable(stock=50)
    post_order(client, order_payload([("creme-test", 1)], name="Salma Benali", phone="0612345678"))
    post_order(client, order_payload([("creme-test", 2)], name="Youssef Alaoui", phone="0700112233", city="Rabat"))
    return Order.objects.order_by("id")


@pytest.mark.django_db
def test_liste_filtres_recherche(staff, orders):
    api = staff(M)
    first, second = orders
    assert api.get("/api/v1/admin/orders").json()["total"] == 2
    assert [o["number"] for o in api.get("/api/v1/admin/orders?search=youssef").json()["items"]] == [second.number]
    assert api.get("/api/v1/admin/orders?search=0612345678").json()["items"][0]["number"] == first.number
    assert api.get(f"/api/v1/admin/orders?search={second.number}").json()["total"] == 1
    assert api.get("/api/v1/admin/orders?city=rabat").json()["total"] == 1
    api.post(f"/api/v1/admin/orders/{first.pk}/status", {"status": "confirmed"}, format="json")
    assert api.get("/api/v1/admin/orders?status=confirmed").json()["items"][0]["number"] == first.number
    assert api.get("/api/v1/admin/orders?from=2000-01-01&to=2000-01-02").json()["total"] == 0
    assert api.get("/api/v1/admin/orders?from=pas-une-date").status_code == 400


@pytest.mark.django_db
@pytest.mark.parametrize("role,sees_cost", [(OWN, True), (A, True), (N, True), (M, False)])
def test_detail_et_cout_selon_role(staff, orders, role, sees_cost):
    data = staff(role).get(f"/api/v1/admin/orders/{orders[0].pk}").json()
    assert data["address"] == "12 rue des Orangers, Maarif" and data["allowedTransitions"] == ["cancelled", "confirmed"]
    assert ("unitCost" in data["items"][0]) is sees_cost


@pytest.mark.django_db
def test_note_interne(staff, orders):
    api = staff(M)
    response = api.patch(f"/api/v1/admin/orders/{orders[0].pk}", {"internalNote": "Appeler avant 18 h"}, format="json")
    assert response.json()["internalNote"] == "Appeler avant 18 h"


@pytest.mark.django_db
def test_facture_imprimable_echappe_les_donnees(client, staff, sellable):
    sellable(stock=5)
    post_order(client, order_payload([("creme-test", 1)], name='<script>alert("x")</script>'))
    order = Order.objects.get()
    response = staff(M).get(f"/api/v1/admin/orders/{order.pk}/invoice")
    html = response.content.decode()
    assert response.status_code == 200 and response["Content-Type"].startswith("text/html")
    assert order.number in html
    assert "149.00 DH" in html or "149,00 DH" in html  # total = 129 + 20 de livraison
    assert "<script>alert" not in html and "&lt;script&gt;" in html
    assert "Paiement à la livraison" in html


@pytest.mark.django_db
def test_export_excel_et_injection_de_formule(client, staff, sellable):
    sellable(stock=5)
    post_order(client, order_payload([("creme-test", 1)], name='=HYPERLINK("http://pirate","x")'))
    response = staff(N).get("/api/v1/admin/orders/export.xlsx")
    assert response.status_code == 200 and "spreadsheetml" in response["Content-Type"]
    ws = load_workbook(io.BytesIO(response.content)).active
    header = [c.value for c in ws[1]]
    row = dict(zip(header, [c.value for c in ws[2]], strict=True))
    assert row["N° commande"] == Order.objects.get().number and row["Total (DH)"] == 149
    assert row["Client"].startswith("'=")  # neutralisé : Excel ne l'exécutera pas
    assert ws.cell(row=2, column=5).data_type == "s"


@pytest.mark.django_db
def test_clients_liste_et_historique(staff, orders):
    api = staff(M)
    listing = api.get("/api/v1/admin/customers?search=salma").json()["items"]
    assert [c["fullName"] for c in listing] == ["Salma Benali"]
    detail = api.get(f"/api/v1/admin/customers/{listing[0]['id']}").json()
    assert detail["phone"] == "+212612345678" and detail["ordersCount"] == 1
    assert [o["number"] for o in detail["orders"]] == [orders[0].number]
    assert detail["addresses"][0]["city"] == "Casablanca"


@pytest.mark.django_db
@pytest.mark.parametrize("role", [*Role, None])
def test_matrice_des_permissions(staff, orders, shop, role):
    api = staff(role)
    order = orders[0]
    read_ops = {OWN, A, M, N}
    write_ops = {OWN, A, M}
    settings_write = {OWN, A}

    def expect(allowed, ok):
        return ok if role in allowed else 403

    assert api.get("/api/v1/admin/orders").status_code == expect(read_ops, 200)
    assert api.get(f"/api/v1/admin/orders/{order.pk}/invoice").status_code == expect(read_ops, 200)
    assert api.get("/api/v1/admin/orders/export.xlsx").status_code == expect(read_ops, 200)
    assert api.get("/api/v1/admin/customers").status_code == expect(read_ops, 200)
    assert api.post(
        f"/api/v1/admin/orders/{order.pk}/status", {"status": "confirmed"}, format="json"
    ).status_code == expect(write_ops, 200)
    assert api.get("/api/v1/admin/coupons").status_code == expect(read_ops, 200)
    assert api.post(
        "/api/v1/admin/coupons",
        {"code": f"C{role}"[:20].replace(".", ""), "type": "percentage", "value": "5"},
        format="json",
    ).status_code == expect(settings_write, 201)
    assert api.patch(f"/api/v1/admin/shipping/zones/{shop.pk}", {"fee": "25"}, format="json").status_code == expect(
        settings_write, 200
    )
    assert api.patch("/api/v1/admin/settings", {"whatsappNumber": "212600000000"}, format="json").status_code == expect(
        settings_write, 200
    )


@pytest.mark.django_db
def test_codes_promo_admin(staff, client, sellable):
    api = staff()
    bad = api.post("/api/v1/admin/coupons", {"code": "TROP", "type": "percentage", "value": "150"}, format="json")
    assert bad.status_code == 400
    created = api.post(
        "/api/v1/admin/coupons", {"code": "bienvenue", "type": "fixed_amount", "value": "30"}, format="json"
    ).json()
    assert created["code"] == "BIENVENUE" and created["timesUsed"] == 0
    sellable(stock=5)
    post_order(client, order_payload([("creme-test", 1)], promo="BIENVENUE"))
    assert api.get(f"/api/v1/admin/coupons/{created['id']}").json()["timesUsed"] == 1
    used = api.delete(f"/api/v1/admin/coupons/{created['id']}")
    assert used.status_code == 409 and used.json()["error"]["code"] == "coupon_in_use"


@pytest.mark.django_db
def test_config_publique_conforme_au_front(client, db):
    """Après import, GET /store/config reproduit src/config/store.js — sans jamais les codes promo."""
    call_command("import_store_config", stdout=open("/dev/null", "w"))  # noqa: SIM115
    front = json.loads(
        (Path(__file__).resolve().parents[1] / "apps/store_settings/fixtures/demo_store_config.json").read_text(
            encoding="utf-8"
        )
    )
    data = client.get("/api/v1/store/config").json()
    assert data["cities"] == front["cities"]
    assert data["deliveryFee"] == front["deliveryFee"]
    assert data["freeDeliveryThreshold"] == front["freeDeliveryThreshold"]
    assert data["whatsappNumber"] == front["whatsappNumber"] and data["allowOtherCity"] is True
    assert data["maxQtyPerLine"] == front["maxQtyPerLine"]
    assert "NATURE10" not in json.dumps(data) and "promo" not in json.dumps(data).lower()


@pytest.mark.django_db
def test_reglages_et_whatsapp(staff, client, shop):
    api = staff(OWN)
    assert api.patch("/api/v1/admin/settings", {"whatsappNumber": "+212 6"}, format="json").status_code == 400
    api.patch("/api/v1/admin/settings", {"whatsappNumber": "212612345678", "ice": "001234567000089"}, format="json")
    assert client.get("/api/v1/store/config").json()["whatsappNumber"] == "212612345678"
