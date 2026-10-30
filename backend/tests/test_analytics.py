"""
Phase 5 — chaque indicateur du tableau de bord est comparé à une requête SQL écrite à la main (MySQL),
indépendante du code testé : bornes de dates calculées ici avec zoneinfo, pas avec apps.analytics.periods.
"""

import datetime as dt
import io
from decimal import Decimal
from zoneinfo import ZoneInfo

import pytest
from django.db import connection
from openpyxl import load_workbook

from apps.accounts.roles import Role
from apps.analytics.periods import Period, resolve
from apps.catalog.models import Category, ProductCategory, Variant
from apps.orders.models import Order
from tests.conftest import order_payload, post_order

URL = "/api/v1/admin/analytics/report"
EXPORT = "/api/v1/admin/analytics/export"
CASA = ZoneInfo("Africa/Casablanca")
CURRENT = {"preset": "custom", "from": "2026-04-01", "to": "2026-04-10"}  # comparaison : 22 → 31 mars


def local(text):
    """« 2026-04-01 00:30 » à l'heure du Maroc → instant conscient (UTC+1 en avril 2026)."""
    return dt.datetime.fromisoformat(text).replace(tzinfo=CASA)


def utc_bounds(first, last):
    """Bornes en UTC « naïf » : c'est ainsi que MySQL stocke les dates (USE_TZ), sans conversion implicite."""
    start, end = local(f"{first} 00:00"), local(f"{last} 00:00") + dt.timedelta(days=1)
    return start.astimezone(dt.UTC).replace(tzinfo=None), end.astimezone(dt.UTC).replace(tzinfo=None)


def sql(query, params=()):
    with connection.cursor() as cursor:
        cursor.execute(query, params)
        return cursor.fetchall()


def one(query, params=()):
    return sql(query, params)[0]


@pytest.fixture
def dataset(client, staff, sellable, category):
    """
    Commandes datées à la main. Clients : A, B, C, D, E, F, G (un téléphone chacun).
    Période étudiée : 1 → 10 avril 2026 ; comparaison : 22 → 31 mars 2026.
    """
    p1 = sellable(slug="serum", stock=500, price="129.00")
    sellable(slug="creme", stock=500, price="99.00")
    # p1 a DEUX catégories : ses ventes ne doivent être comptées qu'une fois (catégorie principale).
    ProductCategory.objects.create(product=p1, category=category(slug="bonus", name="Bonus"), position=5)
    api = staff()
    phones = {c: f"06222222{i:02d}" for i, c in enumerate("ABCDEFG")}

    def order(customer, items, placed, flow=(), **stamps):
        response = post_order(client, order_payload(items, phone=phones[customer], name=f"Client {customer}"))
        assert response.status_code == 201, response.content
        o = Order.objects.get(number=response.json()["number"])
        for status in flow:
            r = api.post(f"/api/v1/admin/orders/{o.pk}/status", {"status": status}, format="json")
            assert r.status_code == 200, r.content
        Order.objects.filter(pk=o.pk).update(placed_at=local(placed), **{k: local(v) for k, v in stamps.items()})
        return o

    delivered = ("confirmed", "shipped", "delivered")
    # --- période précédente (22 → 31 mars)
    order("A", [("serum", 1)], "2026-03-23 10:00", delivered, delivered_at="2026-03-25 15:00")
    order("B", [("creme", 1)], "2026-03-26 09:00", ("cancelled",), cancelled_at="2026-03-26 10:00")
    # passée en mars, livrée en avril : commande de mars, encaissement d'avril
    order("C", [("serum", 2)], "2026-03-31 18:00", delivered, delivered_at="2026-04-02 11:00")
    # --- période étudiée (1 → 10 avril)
    order("A", [("serum", 2)], "2026-04-01 00:30")  # minuit passé à Casablanca = 31 mars 23:30 UTC
    order("B", [("creme", 1)], "2026-04-03 12:00", ("confirmed",))
    order("D", [("serum", 2), ("creme", 1)], "2026-04-05 09:00", delivered, delivered_at="2026-04-06 17:00")
    order(
        "D",
        [("creme", 4)],
        "2026-04-06 09:00",
        (*delivered, "returned"),
        delivered_at="2026-04-07 10:00",
        returned_at="2026-04-08 10:00",
    )
    order("E", [("serum", 1)], "2026-04-07 08:00", ("cancelled",))
    order("F", [("creme", 1)], "2026-04-10 23:45", ("confirmed", "shipped"))  # 22:45 UTC : reste le 10
    order("G", [("serum", 1)], "2026-04-11 00:10")  # 10 avril 23:10 UTC, mais le 11 au Maroc : EXCLUE
    # Les prix du catalogue changent APRÈS les commandes : les chiffres ne doivent pas bouger.
    Variant.objects.update(price=Decimal("999.00"))
    return api


def test_periodes_equivalentes():
    t = dt.date(2026, 3, 31)
    assert resolve("today", t) == (Period(t, t), Period(dt.date(2026, 3, 30), dt.date(2026, 3, 30)))
    assert resolve("7d", t) == (Period(dt.date(2026, 3, 25), t), Period(dt.date(2026, 3, 18), dt.date(2026, 3, 24)))
    assert resolve("30d", t) == (Period(dt.date(2026, 3, 2), t), Period(dt.date(2026, 1, 31), dt.date(2026, 3, 1)))
    assert resolve("90d", t)[0].days == 90 and resolve("90d", t)[1].days == 90
    # Ce mois (1 → 31 mars) vs février 1 → 28 (le 31 n'existe pas en février)
    assert resolve("month", t) == (Period(dt.date(2026, 3, 1), t), Period(dt.date(2026, 2, 1), dt.date(2026, 2, 28)))
    # Mois précédent en janvier : décembre vs novembre (changement d'année)
    assert resolve("last_month", dt.date(2026, 1, 15)) == (
        Period(dt.date(2025, 12, 1), dt.date(2025, 12, 31)),
        Period(dt.date(2025, 11, 1), dt.date(2025, 11, 30)),
    )
    # Année en cours un 29 février : comparée au 1er janvier → 28 février de l'année d'avant
    assert resolve("year", dt.date(2028, 2, 29)) == (
        Period(dt.date(2028, 1, 1), dt.date(2028, 2, 29)),
        Period(dt.date(2027, 1, 1), dt.date(2027, 2, 28)),
    )
    custom = resolve("custom", t, dt.date(2026, 4, 1), dt.date(2026, 4, 10))
    assert custom[1] == Period(dt.date(2026, 3, 22), dt.date(2026, 3, 31))


def test_bornes_au_fuseau_du_maroc_y_compris_ramadan():
    # Pendant le ramadan 2026, le Maroc est à UTC+0 ; ensuite à UTC+1 : les bornes suivent.
    for day in ("2026-03-01", "2026-04-05"):
        d = dt.date.fromisoformat(day)
        start, end = Period(d, d).bounds()
        expected = dt.datetime.combine(d, dt.time.min).replace(tzinfo=CASA)
        assert start == expected and end == expected + dt.timedelta(days=1)
    assert Period(dt.date(2026, 3, 1), dt.date(2026, 3, 1)).bounds()[0].utcoffset() == dt.timedelta(0)
    assert Period(dt.date(2026, 4, 5), dt.date(2026, 4, 5)).bounds()[0].utcoffset() == dt.timedelta(hours=1)


@pytest.mark.django_db
def test_indicateurs_egaux_aux_requetes_sql(dataset):
    data = dataset.get(URL, CURRENT).json()
    assert data["period"] == {"from": "2026-04-01", "to": "2026-04-10", "days": 10}
    assert data["previousPeriod"] == {"from": "2026-03-22", "to": "2026-03-31", "days": 10}

    for key, (first, last) in (("kpis", ("2026-04-01", "2026-04-10")), ("previousKpis", ("2026-03-22", "2026-03-31"))):
        k = data[key]
        start, end = utc_bounds(first, last)
        win = "FROM orders_order WHERE placed_at >= %s AND placed_at < %s"
        orders, valid, gross, cancelled, returned, left = one(
            "SELECT COUNT(*), SUM(status NOT IN ('cancelled','returned')),"
            " COALESCE(SUM(CASE WHEN status NOT IN ('cancelled','returned') THEN total END), 0),"
            " SUM(status = 'cancelled'), SUM(status = 'returned'),"
            " SUM(status IN ('shipped','delivered','returned')) " + win,
            [start, end],
        )
        revenue, delivered = one(
            "SELECT COALESCE(SUM(total), 0), COUNT(*) FROM orders_order"
            " WHERE status = 'delivered' AND delivered_at >= %s AND delivered_at < %s",
            [start, end],
        )
        (units,) = one(
            "SELECT COALESCE(SUM(i.quantity), 0) FROM orders_orderitem i JOIN orders_order o ON o.id = i.order_id"
            " WHERE o.placed_at >= %s AND o.placed_at < %s AND o.status NOT IN ('cancelled','returned')",
            [start, end],
        )
        (customers,) = one(
            "SELECT COUNT(DISTINCT customer_id) " + win + " AND status NOT IN ('cancelled','returned')", [start, end]
        )
        (new,) = one(
            "SELECT COUNT(*) FROM (SELECT customer_id, MIN(placed_at) AS first FROM orders_order"
            " WHERE status NOT IN ('cancelled','returned') GROUP BY customer_id) t WHERE first >= %s AND first < %s",
            [start, end],
        )
        assert k["orders"] == orders and k["validOrders"] == valid
        assert Decimal(str(k["grossSales"])) == gross
        assert Decimal(str(k["revenue"])) == revenue and k["deliveredOrders"] == delivered
        assert k["unitsSold"] == units
        assert k["customers"] == customers and k["newCustomers"] == new and k["returningCustomers"] == customers - new
        assert k["cancelled"] == cancelled and k["returned"] == returned
        assert Decimal(str(k["averageOrderValue"])) == (gross / valid).quantize(Decimal("0.01"))
        assert k["cancellationRate"] == float((Decimal(cancelled) * 100 / orders).quantize(Decimal("0.1")))
        if left:
            assert k["returnRate"] == float((Decimal(returned) * 100 / left).quantize(Decimal("0.1")))

    k = data["kpis"]
    # Valeurs attendues du jeu de données (contrôle humain en plus du SQL)
    assert (k["orders"], k["validOrders"], k["cancelled"], k["returned"]) == (6, 4, 1, 1)  # G exclue (le 11 au Maroc)
    assert k["deliveredOrders"] == 2  # C (passée en mars, livrée le 2 avril) + D ; la retournée n'est pas encaissée
    assert (k["customers"], k["newCustomers"], k["returningCustomers"]) == (
        4,
        3,
        1,
    )  # A existant ; B nouveau malgré sa commande annulée de mars
    assert k["unitsSold"] == 2 + 1 + 3 + 1
    assert (k["cancellationRate"], k["returnRate"]) == (16.7, 33.3)  # 1/6 ; 1 retour / 3 parties chez le client
    assert data["previousKpis"]["orders"] == 3 and data["previousKpis"]["deliveredOrders"] == 1


@pytest.mark.django_db
def test_montants_figes_jamais_les_prix_actuels(dataset):
    k = dataset.get(URL, CURRENT).json()["kpis"]
    (line_value,) = one("SELECT SUM(unit_price * quantity) FROM orders_orderitem WHERE unit_price = 999")
    assert line_value is None  # aucune ligne au nouveau prix
    assert k["grossSales"] < 999  # 4 commandes valides aux anciens prix (129 / 99 DH)


@pytest.mark.django_db
def test_comparaison_avec_la_periode_precedente(dataset):
    data = dataset.get(URL, CURRENT).json()
    k, p, c = data["kpis"], data["previousKpis"], data["changes"]
    assert c["orders"] == {"percent": round((k["orders"] - p["orders"]) * 100 / p["orders"], 1)}
    assert c["revenue"]["percent"] == round((k["revenue"] - p["revenue"]) * 100 / p["revenue"], 1)
    assert c["cancellationRate"] == {"points": round(k["cancellationRate"] - p["cancellationRate"], 1)}
    assert p["returned"] == 0 and c["returned"] is None  # pas de % depuis zéro
    assert p["returnRate"] == 0.0 and c["returnRate"] == {"points": k["returnRate"]}  # A livrée en mars, 0 retour


@pytest.mark.django_db
def test_series_jour_semaine_mois_et_fuseau(dataset):
    for granularity, buckets in (("day", 10), ("week", 2), ("month", 1)):
        data = dataset.get(URL, {**CURRENT, "granularity": granularity}).json()
        s = data["series"]
        assert len(s) == buckets, granularity
        assert s[0]["start"] == "2026-04-01" and s[-1]["end"] == "2026-04-10"
        assert sum(b["orders"] for b in s) == data["kpis"]["orders"]
        assert Decimal(str(sum(b["grossSales"] for b in s))) == Decimal(str(data["kpis"]["grossSales"]))
        assert Decimal(str(sum(b["revenue"] for b in s))) == Decimal(str(data["kpis"]["revenue"]))
    days = {b["start"]: b for b in dataset.get(URL, CURRENT).json()["series"]}
    assert days["2026-04-01"]["orders"] == 1  # 00:30 au Maroc : le 1er, pas le 31 mars
    assert days["2026-04-10"]["orders"] == 1  # 23:45 au Maroc : le 10
    assert days["2026-04-02"]["revenue"] > 0  # encaissement de C au jour de livraison
    week = dataset.get(URL, {**CURRENT, "granularity": "week"}).json()["series"]
    assert (week[0]["start"], week[0]["end"], week[1]["start"]) == ("2026-04-01", "2026-04-05", "2026-04-06")  # lundi 6


@pytest.mark.django_db
def test_repartitions_statut_categorie_paiement(dataset):
    data = dataset.get(URL, CURRENT).json()
    statuses = {s["status"]: s["orders"] for s in data["statuses"]}
    start, end = utc_bounds("2026-04-01", "2026-04-10")
    for status, n in sql(
        "SELECT status, COUNT(*) FROM orders_order WHERE placed_at >= %s AND placed_at < %s GROUP BY status",
        [start, end],
    ):
        assert statuses[status] == n
    assert sum(statuses.values()) == data["kpis"]["orders"]

    expected = dict(
        sql(
            "SELECT c.name, SUM(i.line_total) FROM orders_orderitem i"
            " JOIN orders_order o ON o.id = i.order_id JOIN catalog_variant v ON v.id = i.variant_id"
            " JOIN catalog_productcategory pc ON pc.id = (SELECT pc2.id FROM catalog_productcategory pc2"
            "   WHERE pc2.product_id = v.product_id ORDER BY pc2.position, pc2.id LIMIT 1)"
            " JOIN catalog_category c ON c.id = pc.category_id"
            " WHERE o.placed_at >= %s AND o.placed_at < %s AND o.status NOT IN ('cancelled','returned')"
            " GROUP BY c.name",
            [start, end],
        )
    )
    got = {c["name"]: Decimal(str(c["sales"])) for c in data["categories"]}
    assert got == expected and "Bonus" not in got  # catégorie secondaire : jamais de double comptage
    assert round(sum(c["share"] for c in data["categories"])) == 100
    assert data["payments"] == [
        {"method": "cod", "label": "Paiement à la livraison", "orders": 4, "grossSales": data["kpis"]["grossSales"]}
    ]


@pytest.mark.django_db
def test_produits_variantes_et_moins_performants(dataset, product):
    product(slug="dormant", status="active")  # actif, jamais vendu
    data = dataset.get(URL, CURRENT).json()["products"]
    start, end = utc_bounds("2026-04-01", "2026-04-10")
    rows = sql(
        "SELECT v.product_id, SUM(i.quantity), SUM(i.line_total) FROM orders_orderitem i"
        " JOIN orders_order o ON o.id = i.order_id JOIN catalog_variant v ON v.id = i.variant_id"
        " WHERE o.placed_at >= %s AND o.placed_at < %s AND o.status NOT IN ('cancelled','returned')"
        " GROUP BY v.product_id ORDER BY SUM(i.quantity) DESC",
        [start, end],
    )
    assert [(r["productId"], r["units"], Decimal(str(r["sales"]))) for r in data["topByUnits"]] == [
        (pid, int(u), s) for pid, u, s in rows
    ]
    assert data["topVariants"][0]["units"] == int(rows[0][1])
    assert data["leastPerforming"][0]["name"] == "Marque Crème test" and data["leastPerforming"][0]["units"] == 0
    assert data["leastPerforming"][0]["productId"] not in {r["productId"] for r in data["topByUnits"]}


@pytest.mark.django_db
def test_stock_zero_et_non_renseigne_jamais_confondus(staff, product):
    p = product(slug="stock", status="active")
    v = p.variants.get()
    Variant.objects.filter(pk=v.pk).update(stock_quantity=0)
    Variant.objects.create(product=p, price=10, stock_quantity=None, title="Inconnue", position=1)
    Variant.objects.create(product=p, price=10, stock_quantity=2, low_stock_threshold=5, title="Basse", position=2)
    Variant.objects.create(product=p, price=10, stock_quantity=0, allow_backorder=True, title="Découvert", position=3)
    archived = product(slug="archive", status="archived")
    Variant.objects.filter(product=archived).update(stock_quantity=None)
    stock = staff().get(URL, CURRENT).json()["stock"]
    assert stock["counts"] == {"low": 1, "outOfStock": 1, "unknown": 1}
    assert [r["variantTitle"] for r in stock["outOfStock"]] == ["Par défaut"] and stock["outOfStock"][0][
        "stockQuantity"
    ] == 0
    assert [r["variantTitle"] for r in stock["unknown"]] == ["Inconnue"] and stock["unknown"][0][
        "stockQuantity"
    ] is None
    assert [r["variantTitle"] for r in stock["low"]] == ["Basse"]


@pytest.mark.django_db
def test_clients_les_plus_actifs_et_villes(dataset):
    data = dataset.get(URL, CURRENT).json()["customers"]
    start, end = utc_bounds("2026-04-01", "2026-04-10")
    top = sql(
        "SELECT c.full_name, COUNT(*), SUM(o.total) FROM orders_order o"
        " JOIN customers_customer c ON c.id = o.customer_id"
        " WHERE o.placed_at >= %s AND o.placed_at < %s AND o.status NOT IN ('cancelled','returned')"
        " GROUP BY c.id, c.full_name ORDER BY COUNT(*) DESC, SUM(o.total) DESC, c.id",
        [start, end],
    )
    assert [(r["name"], r["orders"], Decimal(str(r["spent"]))) for r in data["top"]] == [(n, c, s) for n, c, s in top]
    assert data["cities"] == [
        {"city": "Casablanca", "orders": 4, "grossSales": data["cities"][0]["grossSales"], "share": 100.0}
    ]


@pytest.mark.django_db
def test_periode_vide_sans_donnees_inventees(staff):
    data = staff().get(URL, CURRENT).json()
    k = data["kpis"]
    assert (k["orders"], k["revenue"], k["unitsSold"], k["customers"]) == (0, 0, 0, 0)
    assert k["averageOrderValue"] is None and k["cancellationRate"] is None and k["returnRate"] is None
    assert data["changes"]["orders"] is None  # 0 → 0 : pas d'évolution inventée
    assert data["categories"] == [] and data["payments"] == [] and data["customers"]["top"] == []
    assert len(data["series"]) == 10 and all(b["orders"] == 0 for b in data["series"])


@pytest.mark.django_db
@pytest.mark.parametrize(
    "params,field",
    [
        ({"preset": "hier"}, "preset"),
        ({"preset": "custom", "from": "2026-04-01"}, "from"),
        ({"preset": "custom", "from": "2026-04-10", "to": "2026-04-01"}, "from"),
        ({"preset": "custom", "from": "2024-01-01", "to": "2026-01-01"}, "from"),
        ({"preset": "custom", "from": "2026-13-01", "to": "2026-12-01"}, "from"),
        ({"preset": "30d", "granularity": "heure"}, "granularity"),
    ],
)
def test_parametres_invalides(staff, params, field):
    response = staff().get(URL, params)
    assert response.status_code == 400 and field in response.json()["error"]["fields"]


@pytest.mark.django_db
@pytest.mark.parametrize("preset", ["today", "7d", "30d", "90d", "month", "last_month", "year"])
def test_tous_les_filtres(staff, preset):
    data = staff().get(URL, {"preset": preset}).json()
    assert data["preset"] == preset and data["period"]["days"] >= 1 and data["previousPeriod"]["days"] >= 1


@pytest.mark.django_db
@pytest.mark.parametrize(
    "role,code",
    [
        (Role.OWNER, 200),
        (Role.ADMIN, 200),
        (Role.ORDER_MANAGER, 200),
        (Role.ANALYST, 200),
        (Role.CONTENT_EDITOR, 403),
        (None, 403),
    ],
)
def test_permissions(staff, role, code):
    api = staff(role)
    assert api.get(URL, CURRENT).status_code == code
    assert api.get(EXPORT, CURRENT).status_code == code


@pytest.mark.django_db
def test_jamais_de_cout_ni_de_telephone(dataset):
    raw = dataset.get(URL, CURRENT).content.decode()
    assert "cost" not in raw.lower() and "40.0" not in raw  # coût d'achat des fixtures : 40 DH
    assert "0622222" not in raw and "+21262222" not in raw and "Orangers" not in raw  # ni téléphone ni adresse


@pytest.mark.django_db
def test_export_excel_identique_a_l_ecran(dataset):
    Order.objects.filter(ship_full_name="Client D").update(ship_full_name="=HYPERLINK(1)")
    from apps.customers.models import Customer

    Customer.objects.filter(full_name="Client D").update(full_name="=HYPERLINK(1)")
    data = dataset.get(URL, CURRENT).json()
    response = dataset.get(EXPORT, CURRENT)
    assert response.status_code == 200
    assert response["Content-Disposition"] == 'attachment; filename="rapport-2026-04-01-2026-04-10.xlsx"'
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == [
        "Synthèse",
        "Évolution",
        "Statuts",
        "Catégories",
        "Paiement",
        "Produits",
        "Variantes",
        "Stock",
        "Clients",
        "Villes",
    ]
    summary = {row[0]: row[1] for row in wb["Synthèse"].iter_rows(min_row=2, values_only=True) if row and row[0]}
    assert summary["Commandes passées (tous statuts)"] == data["kpis"]["orders"]
    assert summary["Chiffre d'affaires encaissé (DH) — livrées, à la date de livraison"] == data["kpis"]["revenue"]
    names = [row[0] for row in wb["Clients"].iter_rows(min_row=2, values_only=True)]
    assert "'=HYPERLINK(1)" in names and "=HYPERLINK(1)" not in names  # injection de formule neutralisée


@pytest.mark.django_db
def test_categorie_supprimee_du_produit_reste_comptee(dataset):
    # Un produit sans catégorie reste compté (« Sans catégorie ») : aucune vente ne disparaît.
    ProductCategory.objects.filter(product__slug="creme").delete()
    data = dataset.get(URL, CURRENT).json()
    names = [c["name"] for c in data["categories"]]
    assert "Sans catégorie" in names
    total = sum(Decimal(str(c["sales"])) for c in data["categories"])
    (expected,) = one(
        "SELECT SUM(i.line_total) FROM orders_orderitem i JOIN orders_order o ON o.id = i.order_id"
        " WHERE o.placed_at >= %s AND o.placed_at < %s AND o.status NOT IN ('cancelled','returned')",
        list(utc_bounds("2026-04-01", "2026-04-10")),
    )
    assert total == expected
    assert Category.objects.filter(slug="cat-creme").exists()
