"""Export Excel du rapport : une feuille par analyse, montants en DH, taux en %. Mêmes chiffres que l'écran."""

import io

from openpyxl import Workbook
from openpyxl.styles import Font

from apps.orders.exports import safe_text

KPI_LABELS = [
    ("revenue", "Chiffre d'affaires encaissé (DH) — livrées, à la date de livraison"),
    ("orders", "Commandes passées (tous statuts)"),
    ("validOrders", "Commandes valides (hors annulées et retournées)"),
    ("grossSales", "Ventes commandées (DH) — commandes valides"),
    ("averageOrderValue", "Panier moyen (DH)"),
    ("unitsSold", "Unités vendues"),
    ("customers", "Clients (au moins une commande valide)"),
    ("newCustomers", "Nouveaux clients"),
    ("returningCustomers", "Clients existants"),
    ("ordersPerCustomer", "Commandes par client"),
    ("valuePerCustomer", "Valeur moyenne par client (DH)"),
    ("cancellationRate", "Taux d'annulation (%)"),
    ("returnRate", "Taux de retour (%)"),
]


def _sheet(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([safe_text(v) if isinstance(v, str) else v for v in row])
    for i, width in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = width
    ws.freeze_panes = "A2"
    return ws


def _change(change):
    if not change:
        return None
    return change.get("percent", change.get("points"))


def report_workbook(data: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    p, q = data["period"], data["previousPeriod"]
    summary = _sheet(
        wb,
        "Synthèse",
        ["Indicateur", f"{p['from']} → {p['to']}", f"Comparaison {q['from']} → {q['to']}", "Évolution (% ou points)"],
        [[label, data["kpis"][k], data["previousKpis"][k], _change(data["changes"][k])] for k, label in KPI_LABELS],
        (62, 22, 30, 24),
    )
    summary.append([])
    summary.append(["Montants figés dans les commandes ; heure du Maroc. Définitions : docs/ARCHITECTURE.md."])

    _sheet(
        wb,
        "Évolution",
        ["Du", "Au", "Commandes", "Ventes commandées (DH)", "Encaissé (DH)"],
        [[s["start"], s["end"], s["orders"], s["grossSales"], s["revenue"]] for s in data["series"]],
        (12, 12, 12, 22, 16),
    )
    _sheet(wb, "Statuts", ["Statut", "Commandes"], [[s["label"], s["orders"]] for s in data["statuses"]], (28, 12))
    _sheet(
        wb,
        "Catégories",
        ["Catégorie (principale)", "Unités", "Ventes (DH)", "Part (%)"],
        [[c["name"], c["units"], c["sales"], c["share"]] for c in data["categories"]],
        (30, 10, 14, 10),
    )
    _sheet(
        wb,
        "Paiement",
        ["Mode de paiement", "Commandes valides", "Ventes (DH)"],
        [[x["label"], x["orders"], x["grossSales"]] for x in data["payments"]],
        (28, 18, 14),
    )
    pr = data["products"]
    _sheet(
        wb,
        "Produits",
        ["Classement", "Produit", "Unités", "Commandes", "Ventes (DH)"],
        [["Plus vendus (unités)", r["name"], r["units"], r["orders"], r["sales"]] for r in pr["topByUnits"]]
        + [["Plus gros CA", r["name"], r["units"], r["orders"], r["sales"]] for r in pr["topBySales"]]
        + [["Moins performants", r["name"], r["units"], r["orders"], r["sales"]] for r in pr["leastPerforming"]],
        (22, 40, 10, 12, 14),
    )
    _sheet(
        wb,
        "Variantes",
        ["Produit", "Variante", "SKU", "Unités", "Ventes (DH)"],
        [[r["name"], r["variantTitle"], r["sku"], r["units"], r["sales"]] for r in pr["topVariants"]],
        (40, 20, 16, 10, 14),
    )
    st = data["stock"]
    _sheet(
        wb,
        "Stock",
        ["État", "Produit", "Variante", "Stock", "Seuil d'alerte"],
        [
            ["Stock bas", r["productName"], r["variantTitle"], r["stockQuantity"], r["lowStockThreshold"]]
            for r in st["low"]
        ]
        + [
            ["Épuisé (0)", r["productName"], r["variantTitle"], r["stockQuantity"], r["lowStockThreshold"]]
            for r in st["outOfStock"]
        ]
        + [
            ["Non renseigné", r["productName"], r["variantTitle"], "non renseigné", r["lowStockThreshold"]]
            for r in st["unknown"]
        ],
        (16, 40, 20, 14, 14),
    )
    cu = data["customers"]
    _sheet(
        wb,
        "Clients",
        ["Client", "Commandes valides", "Montant (DH)"],
        [[r["name"], r["orders"], r["spent"]] for r in cu["top"]],
        (34, 18, 14),
    )
    _sheet(
        wb,
        "Villes",
        ["Ville", "Commandes valides", "Ventes (DH)", "Part (%)"],
        [[c["city"], c["orders"], c["grossSales"], c["share"]] for c in cu["cities"]],
        (24, 18, 14, 10),
    )
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
